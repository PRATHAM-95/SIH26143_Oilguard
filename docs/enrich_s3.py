"""Rebuild slide 3 as pictorial 3-swimlane layout.
Keeps template chrome (GROUP shapes); wipes everything else added earlier,
then embeds rendered illustration PNGs + detail cards + tech chips.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

PPTX = Path(r"C:\Users\yoges\OneDrive\Desktop\SIH26143_OIL_DUMP\SIH-Oil-Spill-Idea.pptx")
A = Path(r"C:\Users\yoges\OneDrive\Desktop\SIH26143_OIL_DUMP\docs\assets")

C = dict(navy="1E40AF", navy_d="1E3A8A", blue_e="DBEAFE",
         teal="0F766E", teal_d="115E59", teal_e="CCFBF1",
         purple="6D28D9", purple_d="5B21B6", purple_e="EDE9FE",
         amber="B45309", amber_d="92400E", amber_e="FEF3C7",
         green="166534", green_d="14532D", green_e="DCFCE7",
         slate="334155", gray="64748B", gray_f="F1F5F9", white="FFFFFF")
def rgb(k): return RGBColor.from_string(C[k])
def IN(v): return Emu(int(v * 914400))

def remove(sh):
    el = sh._element
    el.getparent().remove(el)

prs = Presentation(str(PPTX))
s = prs.slides[2]
kept = removed = 0
for sh in list(s.shapes):
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        kept += 1
    else:
        remove(sh); removed += 1
print(f"S3 wipe: kept {kept} chrome groups, removed {removed} shapes")

def style_tf(tf, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, size, bold, color, sb) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if sb: p.space_before = Pt(sb)
        r = p.add_run(); r.text = txt
        f = r.font; f.size = Pt(size); f.bold = bold
        f.color.rgb = rgb(color); f.name = "Segoe UI"

def add_card(l, t, w, h, fl, ln, lines):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(l), IN(t), IN(w), IN(h))
    try: sp.adjustments[0] = 0.08
    except Exception: pass
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fl)
    sp.line.color.rgb = rgb(ln); sp.line.width = Pt(1.75)
    sp.shadow.inherit = False
    style_tf(sp.text_frame, lines)

def add_text(l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(IN(l), IN(t), IN(w), IN(h))
    style_tf(tb.text_frame, lines, align, anchor)

def arrow(cx, cy, ch="\u25b6"):
    add_text(cx - 0.30, cy - 0.26, 0.60, 0.52,
             [(ch, 20, True, "gray", 0)], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

def card_lines(title, subs, tc):
    return [(title, 14.5, True, tc, 0)] + [(x, 9.5, False, "slate", 3) for x in subs]

LANES = [1.62, 4.32, 7.02]
LH = 2.45

# ---- Lane labels ----
for ly, lab in zip(LANES, [
        "STAGE A \u00b7 DETECT \u2014 FIND THE SLICK",
        "STAGE B \u00b7 REWIND \u2014 OCEAN PHYSICS IN REVERSE",
        "STAGE C \u00b7 ATTRIBUTE \u2192 PROVE \u2014 NAME THE SHIP"]):
    add_text(0.57, ly - 0.30, 12.0, 0.30, [(lab, 11, True, "gray", 0)])

# ---- Lane 1 : DETECT ----
y = LANES[0]
s.shapes.add_picture(str(A / "s3_sat_sar.png"), IN(0.55), IN(y), IN(4.10), IN(LH))
add_card(4.95, y, 4.00, LH, "blue_e", "navy",
    card_lines("PREPROCESS", ["IW GRD VV+VH \u00b7 10 m resolution",
                              "orbit revisit 6\u201312 days",
                              "Lee speckle filter + land mask"], "navy_d"))
add_card(9.35, y, 5.00, LH, "teal_e", "teal",
    card_lines("U-NET SEGMENTATION",
               ["trained on Zenodo Parts I\u2013III (~3,500 scenes)",
                "classes: slick \u00b7 ship wake \u00b7 biogenic \u00b7 rain \u00b7 low-wind",
                "dual-pol features kill look-alikes"], "teal_d"))
add_card(14.75, y, 4.70, LH, "teal_e", "teal",
    card_lines("SLICK POLYGONS",
               ["GeoJSON polygon + confidence mask",
                "area \u00b7 perimeter \u00b7 shape index",
                "backscatter statistics attached"], "teal_d"))
arrow(4.80, y + LH / 2); arrow(9.15, y + LH / 2); arrow(14.55, y + LH / 2)

# ---- Lane 2 : REWIND ----
y = LANES[1]
s.shapes.add_picture(str(A / "s3_drift_rewind.png"), IN(0.55), IN(y), IN(4.10), IN(LH))
add_card(4.95, y, 5.00, LH, "purple_e", "purple",
    card_lines("OPEN DRIFT \u00b7 BACKWARD HINDCAST",
               ["OpenOil ensemble \u00d7100 virtual particles",
                "forcing: CMEMS currents + ERA5 winds",
                "wind drift 1\u20136% \u00b7 Ekman angle \u00b7 Stokes drift"], "purple_d"))
add_card(10.35, y, 4.00, LH, "purple_e", "purple",
    card_lines("ORIGIN PROBABILITY MAP",
               ["PDF over sea surface \u2014 never a single guess",
                "hindcast window 2\u201312 h",
                "uncertainty reported, not hidden"], "purple_d"))
s.shapes.add_picture(str(A / "s3_heatmap_mini.png"), IN(14.75), IN(y), IN(4.70), IN(LH))
arrow(4.80, y + LH / 2); arrow(10.15, y + LH / 2); arrow(14.55, y + LH / 2)

# ---- Lane 3 : ATTRIBUTE -> PROVE ----
y = LANES[2]
s.shapes.add_picture(str(A / "s3_ais_fusion.png"), IN(0.55), IN(y), IN(4.10), IN(LH))
add_card(4.95, y, 5.00, LH, "amber_e", "amber",
    card_lines("AIS FUSION SCORING",
               ["proximity to origin mode .40",
                "trajectory bearing match .25 \u00b7 anomaly .15",
                "AIS dark-window during spill .20"], "amber_d"))
s.shapes.add_picture(str(A / "s3_dossier_mock.png"), IN(10.35), IN(y), IN(4.00), IN(LH))
add_card(14.75, y, 4.70, LH, "green_e", "green",
    card_lines("OUTPUTS",
               ["dashboard map overlay \u00b7 ranked suspects",
                "court-ready PDF case-file + GeoJSON export",
                "< 5 min inference per scene, Colab-class GPU"], "green_d"))
arrow(4.80, y + LH / 2); arrow(10.15, y + LH / 2); arrow(14.55, y + LH / 2)

# ---- Tech chips ----
chips = ["PyTorch \u00b7 smp", "FastAPI \u00b7 MongoDB", "React \u00b7 Leaflet",
         "OpenDrift", "CMEMS \u00b7 ERA5", "MarineCadastre AIS"]
for i, chp in enumerate(chips):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            IN(0.87 + i * 3.055), IN(9.72), IN(2.92), IN(0.52))
    try: sp.adjustments[0] = 0.5
    except Exception: pass
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb("gray_f")
    sp.line.color.rgb = rgb("gray"); sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    style_tf(sp.text_frame, [(chp, 11, True, "slate", 0)])

prs.save(str(PPTX))
print("SAVED:", PPTX.name)
